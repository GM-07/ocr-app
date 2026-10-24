"""Shared builders for tests. The reader is always a fake."""

from __future__ import annotations

from pathlib import Path

import pypdfium2 as pdfium
from docx import Document

from ocr_app.application.extract_text import ExtractText
from ocr_app.application.ports import Reader
from ocr_app.domain.models import ExportFormat, TextBlock
from ocr_app.infrastructure.docx_exporter import DocxExporter
from ocr_app.infrastructure.loaders import DocumentLoader
from ocr_app.infrastructure.pdf_exporter import PdfExporter
from ocr_app.infrastructure.preprocessor import ImagePreprocessor


def _remember(shapes: list[tuple[int, ...]], image) -> None:
    if getattr(image, "ndim", None) != 3 or image.shape[2] != 3:
        raise AssertionError(f"expected an HxWx3 image, got {getattr(image, 'shape', None)}")
    shapes.append(tuple(image.shape))


class FakeReader:
    """Returns the same blocks for every page and records image shapes."""

    def __init__(self, blocks: list[TextBlock]) -> None:
        self._blocks = list(blocks)
        self.shapes: list[tuple[int, ...]] = []

    def read(self, image) -> list[TextBlock]:
        _remember(self.shapes, image)
        return list(self._blocks)


class ScriptedReader:
    """Returns a different block list for each page, in call order."""

    def __init__(self, pages: list[list[TextBlock]]) -> None:
        self._pages = pages
        self.shapes: list[tuple[int, ...]] = []

    def read(self, image) -> list[TextBlock]:
        _remember(self.shapes, image)
        return list(self._pages[len(self.shapes) - 1])


def build_use_case(reader: Reader, threshold: float = 0.4) -> ExtractText:
    return ExtractText(
        reader=reader,
        preprocessor=ImagePreprocessor(),
        loader=DocumentLoader(),
        exporters={
            ExportFormat.PDF: PdfExporter(),
            ExportFormat.DOCX: DocxExporter(),
        },
        confidence_threshold=threshold,
    )


def docx_text(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def pdf_text(path: Path) -> str:
    pdf = pdfium.PdfDocument(str(path))
    try:
        parts: list[str] = []
        for index in range(len(pdf)):
            page = pdf[index]
            text_page = page.get_textpage()
            try:
                parts.append(text_page.get_text_bounded())
            finally:
                text_page.close()
                page.close()
        return "\n".join(parts)
    finally:
        pdf.close()
